from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r"E:\ICSIP2026_PV_Contamination_Paper\Template.docx")
OUTPUT_LOCAL = ROOT / "paper" / "PV_Contamination_ICSIP_Draft.docx"
OUTPUT_EXTERNAL = Path(r"E:\ICSIP2026_PV_Contamination_Paper\PV_Contamination_ICSIP_Draft.docx")
ASSET_DIR = ROOT / "paper" / "assets"

TITLE = "A Robust Lightweight Framework for UAV-Based Photovoltaic Panel Contamination Detection Under Complex Illumination"
ABSTRACT = (
    "Visible-light inspection of photovoltaic plants from unmanned aerial vehicles offers a low-cost alternative "
    "to manual patrol, but contamination detection remains difficult under strong reflection, cast shadow, viewpoint "
    "variation, and mixed background clutter. This paper presents a lightweight panel-aware framework for PV "
    "contamination detection that combines illumination-robust preprocessing, CBAM-enhanced YOLOv8 panel "
    "segmentation, and a YOLO11 detector for dust, crack, and biological stain localization at the panel level. "
    "The framework is designed for a self-built dataset of approximately 3000 UAV and field images, and its "
    "deployment target is an RTX 4060-class inspection workstation. A preliminary dust-only model already achieved "
    "an F1-score of 0.97, supporting the feasibility of the proposed unified framework."
)
KEYWORDS = "photovoltaic inspection; contamination detection; UAV vision; object detection; image processing; complex illumination"


def clear_document(doc: Document) -> None:
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def add_paragraph(doc: Document, text: str, *, bold: bool = False, size: int = 11, center: bool = False) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_heading(doc: Document, text: str) -> None:
    add_paragraph(doc, text, bold=True, size=12)


def add_caption(doc: Document, text: str) -> None:
    add_paragraph(doc, text, size=10, center=True)
    doc.paragraphs[-1].runs[0].italic = True


def add_picture_if_exists(doc: Document, path: Path, width: float) -> None:
    if path.exists():
        doc.add_picture(str(path), width=Inches(width))


def set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.size = Pt(10)


def build_doc() -> Document:
    doc = Document(str(TEMPLATE)) if TEMPLATE.exists() else Document()
    clear_document(doc)

    add_paragraph(doc, TITLE, bold=True, size=16, center=True)
    add_paragraph(doc, "Author information to be completed before submission", size=11, center=True)
    add_paragraph(doc, "Replace this block with the final author list, affiliations, and emails.", size=10, center=True)

    add_heading(doc, "Abstract")
    add_paragraph(doc, ABSTRACT)

    add_heading(doc, "Keywords")
    add_paragraph(doc, KEYWORDS)

    add_heading(doc, "1 Introduction")
    add_paragraph(
        doc,
        "Photovoltaic plants operate in open outdoor environments and therefore suffer from surface contamination "
        "and structural degradation caused by dust deposition, cracks, and biological residues. UAV-based RGB "
        "inspection is attractive because it provides wide-area coverage with relatively low equipment cost, but "
        "strong reflection, cast shadow, viewpoint variation, and background clutter make robust detection difficult.",
    )
    add_paragraph(
        doc,
        "This paper focuses on a single technical line: panel-aware contamination detection and localization under "
        "complex illumination. The method first segments PV panels and then detects contamination only inside "
        "panel-aware regions of interest.",
    )

    add_heading(doc, "2 Related Work")
    add_paragraph(
        doc,
        "Recent PV inspection work spans infrared diagnosis, electroluminescence analysis, and RGB-image-based "
        "contamination analysis. Aerial PV panel detection is also relevant because panel extraction quality directly "
        "affects downstream contamination detection. Attention mechanisms such as CBAM and class-imbalance-aware "
        "losses such as Focal Loss are useful in this setting because the scene is visually complex and the defect "
        "categories are imbalanced.",
    )

    add_heading(doc, "3 Method")
    add_paragraph(
        doc,
        "The proposed framework first applies illumination normalization to suppress overexposure and improve local "
        "contrast in shadow regions. A YOLOv8m-seg model with CBAM-enhanced features then extracts panel masks. "
        "Each segmented panel is cropped and normalized before contamination detection. A YOLO11 detector predicts "
        "three classes: dust, crack, and biological stain. The final analysis module aggregates panel-level coverage "
        "and severity cues into a maintenance-oriented risk score.",
    )
    add_picture_if_exists(doc, ASSET_DIR / "framework_overview.png", 6.3)
    add_caption(doc, "Fig. 1. Overview of the proposed panel-aware contamination detection framework.")

    add_heading(doc, "4 Experiments")
    add_paragraph(
        doc,
        "The project is based on a self-built RGB PV inspection dataset containing approximately 3000 images acquired "
        "from UAV views and close-range field observations. The current draft assumes a 70/15/15 train/validation/test "
        "split; this setting should be updated if the final protocol differs. All experiments target an NVIDIA RTX 4060 GPU.",
    )
    add_paragraph(doc, "Table 1 summarizes the main comparison placeholders that must be filled with real training logs.")

    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = ["Method", "Precision (%)", "Recall (%)", "mAP@0.5 (%)", "mAP@0.5:0.95 (%)", "Params (M)", "FPS"]
    for cell, value in zip(table.rows[0].cells, headers):
        set_cell_text(cell, value, bold=True)
    rows = [
        ["YOLOv8n direct detection", "XX.X", "XX.X", "XX.X", "XX.X", "XX.X", "XX.X"],
        ["YOLOv8m direct detection", "XX.X", "XX.X", "XX.X", "XX.X", "XX.X", "XX.X"],
        ["YOLO11n panel-aware detection", "XX.X", "XX.X", "XX.X", "XX.X", "XX.X", "XX.X"],
        ["Proposed method", "XX.X", "XX.X", "XX.X", "XX.X", "XX.X", "XX.X"],
    ]
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            set_cell_text(cell, value)
    add_caption(doc, "Table 1. Main comparison on the PV contamination test set.")

    add_picture_if_exists(doc, ASSET_DIR / "pv_examples.png", 6.3)
    add_caption(doc, "Fig. 2. Representative contamination patterns in the project dataset.")
    add_picture_if_exists(doc, ASSET_DIR / "qualitative_results.png", 6.3)
    add_caption(doc, "Fig. 3. Qualitative panel-aware detection visualization.")

    add_heading(doc, "5 Conclusion")
    add_paragraph(
        doc,
        "This paper presents a panel-aware visual inspection framework for UAV-based PV contamination detection under "
        "complex illumination. The draft already matches the correct ICSIP writing logic; the remaining work before "
        "submission is to replace all placeholder quantitative values with the final training and evaluation logs.",
    )

    add_heading(doc, "References")
    references = [
        "[1] M. Islam et al., \"Artificial Intelligence in Photovoltaic Fault Identification and Diagnosis: A Systematic Review,\" Energies, 2023.",
        "[2] E. Arnaudo et al., \"A Comparative Evaluation of Deep Learning Techniques for Photovoltaic Panel Detection From Aerial Images,\" IEEE Access, 2023.",
        "[3] R. Cavieres et al., \"Automatic soiling and partial shading assessment on PV modules through RGB images analysis,\" Applied Energy, 2022.",
        "[4] B. I. Evstatiev et al., \"PV Module Soiling Detection Using Visible Spectrum Imaging and Machine Learning,\" Energies, 2024.",
        "[5] S. Mehta et al., \"DeepSolarEye: Power Loss Prediction and Weakly Supervised Soiling Localization via Fully Convolutional Networks for Solar Panels,\" WACV, 2018.",
        "[6] A. Rahman, \"Solar Panel Surface Defect and Dust Detection: Deep Learning Approach,\" Journal of Imaging, 2025.",
        "[7] S. Woo et al., \"CBAM: Convolutional Block Attention Module,\" ECCV, 2018.",
        "[8] T.-Y. Lin et al., \"Focal Loss for Dense Object Detection,\" ICCV, 2017.",
        "[9] S. Jumaboev et al., \"Photovoltaics Plant Fault Detection Using Deep Learning Techniques,\" Remote Sensing, 2022.",
        "[10] J. Redmon et al., \"You Only Look Once: Unified, Real-Time Object Detection,\" CVPR, 2016.",
    ]
    for item in references:
        add_paragraph(doc, item, size=10)

    return doc


def main() -> None:
    doc = build_doc()
    OUTPUT_LOCAL.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT_LOCAL))
    doc.save(str(OUTPUT_EXTERNAL))


if __name__ == "__main__":
    main()
